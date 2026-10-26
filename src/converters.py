import csv
import json
import subprocess
import logging
from tempfile import TemporaryDirectory
from pathlib import Path
from docx2pdf import convert as word_to_pdf
from docx import Document

REGISTRY = {}

def register(source_format, target_format):
    def decorator(function):
        key = (source_format.lower(), target_format.lower())
        REGISTRY[key] = function
        return function

    return decorator

@register(".csv", ".json")
def csv_to_json(input_file_path: Path, output_file_path: Path) -> Path:
    with open(input_file_path, mode='r', newline='', encoding='utf-8') as csvfile:
        data = list(csv.DictReader(csvfile))

    with open(output_file_path, mode='w', encoding='utf-8') as jsonfile:
        json.dump(data, jsonfile, ensure_ascii=False, indent=4)
        
    return output_file_path


@register(".json", ".csv")
def json_to_csv(input_file_path: Path, output_file_path: Path) -> Path:
    with open(input_file_path, 'r', encoding='utf-8') as json_file:
        data = json.load(json_file)

    json_check(data)
    
    headers = list(data[0].keys())
    
    with open(output_file_path, "w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=headers)
        writer.writeheader()
        writer.writerows(data)
        
    return output_file_path

def json_check(data):
            
    if not isinstance(data, list):
        raise ValueError("JSON musi zawierać listę rekordów.")
    
    if not data:
        raise ValueError("JSON nie zawiera rekordów — nie można utworzyć CSV.")        

    if not isinstance(data[0], dict):
        raise ValueError("Każdy rekord JSON musi być obiektem.")


@register(".docx", ".pdf")
def docx_to_pdf(input_file_path: Path, output_file_path: Path) -> Path:
    input_file_path = input_file_path.resolve()
    output_file_path = output_file_path.resolve()

    if not input_file_path.is_file():
        raise FileNotFoundError(f"Nie znaleziono pliku: {input_file_path}")

    if not output_file_path.parent.is_dir():
        raise NotADirectoryError("Katalog docelowy nie istnieje")
    
    return docx_to_pdf_word(input_file_path, output_file_path) 

    
def docx_to_pdf_word(input_file_path: Path, output_file_path: Path) -> Path:
    word_to_pdf(input_file_path, output_file_path)
    return output_file_path


def docx_to_pdf_libre(input_file_path: Path, output_file_path: Path) -> Path:
    return "not yet"


@register(".txt", ".docx")
def txt_to_docx(input_file_path: Path, output_file_path: Path) -> Path:
    doc = Document()

    with open(input_file_path, "r", encoding="utf-8", errors="replace") as file:
        for line in file:
            doc.add_paragraph(line.rstrip())

    doc.save(output_file_path)
    return output_file_path


@register(".txt", ".pdf")
def txt_to_pdf(input_file_path: Path, output_file_path: Path) -> Path:
    with TemporaryDirectory() as temp_dir:
        temp_docx = Path(temp_dir) / "document.docx"

        txt_to_docx(input_file_path, temp_docx)
        docx_to_pdf(temp_docx, output_file_path)

    return output_file_path

@register(".docx", ".txt")
def docx_to_txt(input_file_path: Path, output_file_path: Path) -> Path:
    doc = Document(input_file_path)

    with open(output_file_path, "w", encoding="utf-8") as file:
        for paragraph in doc.paragraphs:
            file.write(paragraph.text + "\n")

    return output_file_path


def convert(input_file, output_file):
    input_path = Path(input_file)
    output_path = Path(output_file)

    key = (input_path.suffix.lower(), output_path.suffix.lower())
    if key not in REGISTRY:
        raise ValueError(
            f"Nieobsługiwana konwersja: {key[0]} -> {key[1]}"
        )

    converter = REGISTRY[key]

    return converter(input_path, output_path)
    
